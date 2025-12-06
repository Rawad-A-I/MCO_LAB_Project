#!/usr/bin/env python3
"""
Script to convert the academic report from Markdown to Word (.docx) format
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

def add_heading_with_style(doc, text, level):
    """Add a heading with proper style"""
    heading = doc.add_heading(text, level=level)
    return heading

def add_code_block(doc, code_text):
    """Add a code block with monospace font"""
    para = doc.add_paragraph()
    run = para.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    para.style = 'No Spacing'
    return para

def add_table_from_text(doc, table_text):
    """Parse and add a table from markdown table text"""
    lines = [line.strip() for line in table_text.strip().split('\n') if line.strip()]
    if len(lines) < 2:
        return
    
    # Parse header
    header = [cell.strip() for cell in lines[0].split('|') if cell.strip()]
    if not header:
        return
    
    # Create table
    table = doc.add_table(rows=1, cols=len(header))
    table.style = 'Light Grid Accent 1'
    
    # Add header
    header_cells = table.rows[0].cells
    for i, cell_text in enumerate(header):
        header_cells[i].text = cell_text
        for paragraph in header_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    # Add data rows (skip separator line)
    for line in lines[2:]:
        cells = [cell.strip() for cell in line.split('|') if cell.strip()]
        if len(cells) == len(header):
            row = table.add_row()
            for i, cell_text in enumerate(cells):
                row.cells[i].text = cell_text

def parse_markdown_to_docx(md_file, docx_file):
    """Convert Markdown file to Word document"""
    doc = Document()
    
    # Set document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Read markdown file
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []
    
    while i < len(lines):
        line = lines[i]
        
        # Handle code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                # End of code block
                if code_block_lines:
                    add_code_block(doc, '\n'.join(code_block_lines))
                code_block_lines = []
                in_code_block = False
            else:
                # Start of code block
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
        
        # Handle tables
        if '|' in line and not line.strip().startswith('#'):
            if not in_table:
                in_table = True
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table and table_lines:
                add_table_from_text(doc, '\n'.join(table_lines))
                table_lines = []
                in_table = False
        
        # Handle headings
        if line.startswith('# '):
            add_heading_with_style(doc, line[2:].strip(), 1)
        elif line.startswith('## '):
            add_heading_with_style(doc, line[3:].strip(), 2)
        elif line.startswith('### '):
            add_heading_with_style(doc, line[4:].strip(), 3)
        elif line.startswith('#### '):
            add_heading_with_style(doc, line[5:].strip(), 4)
        elif line.startswith('##### '):
            add_heading_with_style(doc, line[6:].strip(), 5)
        elif line.startswith('**') and line.endswith('**') and ':' in line:
            # Bold labels like "**Course:**"
            para = doc.add_paragraph()
            run = para.add_run(line)
            run.bold = True
        elif line.strip() == '---':
            # Horizontal rule - add spacing
            doc.add_paragraph()
        elif line.strip() and not line.strip().startswith('*') and not line.strip().startswith('-') and not line.strip().startswith('1.'):
            # Regular paragraph
            # Clean up markdown formatting
            text = line.strip()
            # Remove markdown bold
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            # Remove markdown italic
            text = re.sub(r'\*(.*?)\*', r'\1', text)
            # Remove markdown code
            text = re.sub(r'`(.*?)`', r'\1', text)
            
            if text:
                para = doc.add_paragraph(text)
        
        # Handle lists
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            text = line.strip()[2:].strip()
            # Clean markdown formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'\*(.*?)\*', r'\1', text)
            para = doc.add_paragraph(text, style='List Bullet')
        elif re.match(r'^\d+\.\s', line.strip()):
            text = re.sub(r'^\d+\.\s', '', line.strip())
            # Clean markdown formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'\*(.*?)\*', r'\1', text)
            para = doc.add_paragraph(text, style='List Number')
        
        i += 1
    
    # Handle any remaining code block or table
    if code_block_lines:
        add_code_block(doc, '\n'.join(code_block_lines))
    if table_lines:
        add_table_from_text(doc, '\n'.join(table_lines))
    
    # Save document
    doc.save(docx_file)
    print(f"Successfully created {docx_file}")

if __name__ == '__main__':
    parse_markdown_to_docx('Academic_Report.md', 'Final_Report.docx')

